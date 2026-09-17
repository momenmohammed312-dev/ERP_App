import sqlite3, re, datetime, json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

DB = r'F:\pos_offline_desktop_database.sqlite'
DAILY_FILE = r'G:\flutter\Downloads\مرتبات_العمال_-_مصحح_جاهز_للاستيراد.xlsx'
MONTHLY_FILE = r'G:\flutter\Downloads\attendance_import_filled (1).xlsx'
PERIOD = '2026-08'

def norm(s):
    if s is None:
        return None
    s = str(s).strip()
    s = re.sub(r'[إأآا]', 'ا', s)
    s = re.sub(r'ى', 'ي', s)
    s = re.sub(r'ة', 'ه', s)
    s = re.sub(r'\s+', '', s)
    return s

# backup first
import shutil, pathlib
backup = f"F:\\pos_offline_desktop_database_backup_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.sqlite"
if os.path.exists(DB):
    shutil.copy2(DB, backup)
    print(f"backup -> {backup}")

con = sqlite3.connect(DB)
cur = con.cursor()
cur.execute('select staff_id, name, status, basic_salary from staff_table')
staff_rows = cur.fetchall()
print(f"staff rows {len(staff_rows)}")

from collections import defaultdict
by_norm = defaultdict(list)
for sid, name, status, bsal in staff_rows:
    by_norm[norm(name)].append((sid, name, status, bsal))

DUPLICATE_KEEPERS = {
    'STAFF0096': 'ام سما',
    'STAFF0102': 'ام محمود',
    'STAFF0115': 'عزه',
    'STAFF0057': 'ياسمين',
    'STAFF0066': 'دنيا خميس',
    'STAFF0078': 'ام سالم',
    'STAFF0074': 'عبدالله محمد',
    'STAFF0077': 'علي محمد',
    'STAFF0069': 'حبيبه علي',
    'STAFF0106': 'اسماء حسني',
}
SKIP_NAMES = {'بسملة','محمد محمود','ملك','هبه','صفا محمود','عائشة محمد','محمد عيسى','ام حماده','ام احمد','فرح عماد','احمد عبد الكريم'}

def resolve(name):
    key = norm(name)
    if key in {norm(x) for x in SKIP_NAMES}:
        return None
    cands = by_norm.get(key, [])
    if len(cands) == 1:
        return cands[0][0]
    if len(cands) > 1:
        ids = {c[0] for c in cands}
        keeper_ids = ids & set(DUPLICATE_KEEPERS.keys())
        if len(keeper_ids) == 1:
            return list(keeper_ids)[0]
        return None
    return None

# check files exist
print("DAILY exists", os.path.exists(DAILY_FILE))
print("MONTHLY exists", os.path.exists(MONTHLY_FILE))

wb_daily = openpyxl.load_workbook(DAILY_FILE, data_only=True)
daily_records = defaultdict(list)
unresolved_daily_sheets = []
for sn in wb_daily.sheetnames:
    sid = resolve(sn)
    ws = wb_daily[sn]
    if sid is None:
        unresolved_daily_sheets.append(sn)
        continue
    header_cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'التاريخ':
                header_cells.append((cell.row, cell.column))
    for hr, hc in header_cells:
        r = hr + 1
        while True:
            date_cell = ws.cell(row=r, column=hc).value
            if not isinstance(date_cell, datetime.datetime):
                break
            checkin = ws.cell(row=r, column=hc + 2).value
            checkout = ws.cell(row=r, column=hc + 3).value
            daily_records[sid].append((date_cell.date(), checkin, checkout))
            r += 1

wb_month = openpyxl.load_workbook(MONTHLY_FILE, data_only=True)
ws_m = wb_month['ملخص']
monthly_from_file = {}
unresolved_monthly_rows = []
for row in ws_m.iter_rows(min_row=2, values_only=True):
    name = row[0]
    if not name or 'تعليمات' in str(name):
        continue
    sid = resolve(name)
    if sid is None:
        unresolved_monthly_rows.append(name)
        continue
    late, extra, perm, absence = row[1], row[2], row[3], row[4]
    monthly_from_file[sid] = {
        'late_hours': late if isinstance(late, (int, float)) else None,
        'overtime_hours': extra if isinstance(extra, (int, float)) else None,
        'permission_hours': perm if isinstance(perm, (int, float)) else None,
        'absent_days': absence if isinstance(absence, (int, float)) else None,
        'source_name': name,
    }

target_ids = set(daily_records.keys()) | set(monthly_from_file.keys())
print('Resolved daily-sheet employees   :', len(daily_records))
print('Resolved monthly-file employees  :', len(monthly_from_file))
print('Union target employees           :', len(target_ids))
print()
print('Unresolved/skipped daily sheets  :', unresolved_daily_sheets)
print('Unresolved/skipped monthly rows  :', unresolved_monthly_rows)
print()

cur.execute('select staff_id, name, status from staff_table')
status_map = {sid: (name, status) for sid, name, status, _ in staff_rows}
terminated_in_target = [(sid, status_map[sid]) for sid in target_ids if status_map[sid][1] != 'active']
print('Non-active employees accidentally in target set:', terminated_in_target)

# also print some sample
print("\nSample daily:", list(daily_records.items())[:2])
print("\nSample monthly:", list(monthly_from_file.items())[:2])

with open(r'G:\development\POS-Offline-Desktop-main\inspect_tmp\resolution_debug.json', 'w', encoding='utf-8') as f:
    json.dump({
        'target_ids': sorted(target_ids),
        'daily_only': sorted(set(daily_records) - set(monthly_from_file)),
        'monthly_only': sorted(set(monthly_from_file) - set(daily_records)),
        'both': sorted(set(daily_records) & set(monthly_from_file)),
        'unresolved_daily': unresolved_daily_sheets,
        'unresolved_monthly': unresolved_monthly_rows,
    }, f, ensure_ascii=False, indent=2)

print("\nWrote resolution_debug.json")
# don't commit yet, just dry run
con.close()
print("DRY RUN DONE - no DB changes yet")
