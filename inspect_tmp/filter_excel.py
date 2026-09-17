import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
src = r'G:\flutter\Downloads\attendance_import_filled (1).xlsx'
dst = r'G:\flutter\Downloads\attendance_import_اضافي_اذن_بس.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb['ملخص']
# keep same headers, but zero out تأخير (col B) and غياب (col E)
for r in range(2, ws.max_row+1):
    # if row is instructions or empty name, skip
    name = ws.cell(row=r, column=1).value
    if name is None or str(name).strip()=='' or 'تعليمات' in str(name):
        continue
    # تأخير col2 -> 0
    ws.cell(row=r, column=2).value = 0
    # غياب col5 -> 0
    ws.cell(row=r, column=5).value = 0
    # keep إضافي col3 and إذن col4 as is, فترة col6 as is
# remove example sheet to avoid confusion
if 'مثال يومي - أحمد خليفة' in wb.sheetnames:
    wb.remove(wb['مثال يومي - أحمد خليفة'])
# also ensure header for ملاحظة clarifies
ws.cell(row=1, column=2).value = 'تأخير\n(سيُهمل - 0)'
ws.cell(row=1, column=5).value = 'غياب\n(سيُهمل - 0)'
wb.save(dst)
print('SAVED', dst)
# verify
wb2 = openpyxl.load_workbook(dst, data_only=True)
ws2 = wb2['ملخص']
for r in range(1, min(6, ws2.max_row+1)):
    print([str(ws2.cell(row=r, column=c).value) for c in range(1,7)])
