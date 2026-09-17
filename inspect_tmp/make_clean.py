import openpyxl
from openpyxl.styles import Font
# create clean file
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ملخص"
headers = ["اسم الموظف", "تأخير", "إضافي", "إذن", "غياب", "فترة"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True)
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
data = [
    ["أحمد خليفة", 0, 0, 0, 0, "2026-08"],
    ["أميرة", 0, 0, 2, 0, "2026-08"],
    ["شهد", 0, 0, 9.5, 0, "2026-08"],
    ["ياسمين", 0, 0, 4.5, 0, "2026-08"],
]
for r, row in enumerate(data, 2):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
# also test with original path
path = r"G:\flutter\Downloads\attendance_import_clean.xlsx"
wb.save(path)
print("saved", path)

# test dart parse again via python using excel parser? just check file exists
import os
print(os.path.exists(path))
