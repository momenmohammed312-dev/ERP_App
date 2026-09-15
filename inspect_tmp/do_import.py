import sqlite3, re, datetime, json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

DB = r'F:\pos_offline_desktop_database.sqlite'
DAILY_FILE = r'G:\flutter\Downloads\مرتبات_العمال_-_مصحح_جاهز_للاستيراد.xlsx'
MONTHLY_FILE = r'G:\flutter\Downloads\attendance_import_filled (1).xlsx'
PERIOD = '2026-08'

def norm(s):
    if s is None: return None
    s = str(s).strip()
    s = re.sub(r'[إأآا]', 'ا', s)
    s = re.sub(r'ى', 'ي', s)
    s = re.sub(r'ة', 'ه', s)
    s = re.sub(r'\s+', '', s)
    return s

con = sqlite3.connect(DB)
cur = con.cursor()
cur.execute('select staff_id, name, status, basic_salary from staff_table')
staff_rows = cur.fetchall()
from collections import defaultdict
by_norm = defaultdict(list)
for sid, name, status, bsal in staff_rows:
    by_norm[norm(name)].append((sid, name, status, bsal))

DUPLICATE_KEEPERS = {
    'STAFF0096': 'ام سما',
    'STAFF0102': 'ام محمود',
    'STAFF0115': 'عزه',
    'STAFF0057': 'ياسمين',
    'STAFF0066': 'دنيا خميس',
    'STAFF0078': 'ام سالم',
    'STAFF0074': 'عبدالله محمد',
    'STAFF0077': 'علي محمد',
    'STAFF0069': 'حبيبه علي',
    'STAFF0106': 'اسماء حسني',
}
SKIP_NAMES = {'بسملة','محمد محمود','ملك','هبه','صفا محمود','عائشة محمد','محمد عيسى','ام حماده','ام احمد','فرح عماد','احمد عبد الكريم'}
def resolve(name):
    key = norm(name)
    if key in {norm(x) for x in SKIP_NAMES}:
        return None
    cands = by_norm.get(key, [])
    if len(cands) == 1:
        return cands[0][0]
    if len(cands) > 1:
        ids = {c[0] for c in cands}
        keeper_ids = ids & set(DUPLICATE_KEEPERS.keys())
        if len(keeper_ids) == 1:
            return list(keeper_ids)[0]
        return None
    return None

# parse daily
wb_daily = openpyxl.load_workbook(DAILY_FILE, data_only=True)
daily_records = defaultdict(list)
for sn in wb_daily.sheetnames:
    sid = resolve(sn)
    if sid is None:
        continue
    ws = wb_daily[sn]
    header_cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'التاريخ':
                header_cells.append((cell.row, cell.column))
    for hr, hc in header_cells:
        r = hr + 1
        while True:
            date_cell = ws.cell(row=r, column=hc).value
            if not isinstance(date_cell, datetime.datetime):
                break
            checkin = ws.cell(row=r, column=hc + 2).value
            checkout = ws.cell(row=r, column=hc + 3).value
            # normalize time values
            def to_time(v):
                if v is None or (isinstance(v, str) and v.strip()==''): return None
                if isinstance(v, datetime.time): return v
                if isinstance(v, datetime.datetime): return v.time()
                if isinstance(v, str):
                    v=v.strip()
                    if 'غ' in v: return None
                    # try parse HH:MM
                    try:
                        parts=v.replace('؛',':').replace(';',':').split(':')
                        h=int(parts[0]); m=int(parts[1]) if len(parts)>1 else 0
                        return datetime.time(h,m)
                    except: return None
                return None
            ci = to_time(checkin)
            co = to_time(checkout)
            daily_records[sid].append((date_cell.date(), ci, co))
            r += 1

# parse monthly
wb_month = openpyxl.load_workbook(MONTHLY_FILE, data_only=True)
ws_m = wb_month['ملخص']
monthly_from_file = {}
for row in ws_m.iter_rows(min_row=2, values_only=True):
    name = row[0]
    if not name or 'تعليمات' in str(name):
        continue
    sid = resolve(name)
    if sid is None:
        continue
    late, extra, perm, absence = row[1], row[2], row[3], row[4]
    # period from col 6 or default
    period = row[5] if len(row)>5 and row[5] else PERIOD
    if isinstance(period, datetime.datetime):
        period = period.strftime('%Y-%m')
    period = str(period).strip()
    if not period or period=='None':
        period=PERIOD
    monthly_from_file[sid] = {
        'late_hours': float(late) if isinstance(late, (int,float)) else 0.0,
        'overtime_hours': float(extra) if isinstance(extra, (int,float)) else 0.0,
        'permission_hours': float(perm) if isinstance(perm, (int,float)) else 0.0,
        'absent_days': int(absence) if isinstance(absence, (int,float)) else 0,
        'period': period,
    }

# now write to DB in transaction
con.execute("BEGIN")
try:
    inserted_att = 0
    updated_att = 0
    for sid, recs in daily_records.items():
        for d, ci, co in recs:
            # convert to timestamps for DB (drift stores as int seconds? check)
            # we saw earlier attendance date stored as seconds (e.g., 1785531600)
            # checkInTime stored as int? Let's check existing format: it was int like 1785561480 (seconds)
            # So we store as seconds since epoch
            date_ts = int(datetime.datetime(d.year, d.month, d.day, tzinfo=datetime.timezone.utc).timestamp())
            # check existing
            cur.execute("SELECT id FROM attendance_table WHERE staff_id=? AND date=?", (sid, date_ts))
            existing = cur.fetchone()
            # build times
            ci_ts = int(datetime.datetime(d.year, d.month, d.day, ci.hour, ci.minute, tzinfo=datetime.timezone.utc).timestamp()) if ci else None
            co_ts = int(datetime.datetime(d.year, d.month, d.day, co.hour, co.minute, tzinfo=datetime.timezone.utc).timestamp()) if co else None
            # determine status
            if ci is None and co is None:
                status='absent'
            elif ci is None:
                status='absent'
            else:
                # check late: if after 08:10 (grace 10)
                if ci.hour*60+ci.minute > 8*60+10:
                    status='late'
                else:
                    status='present'
                # if early leave? check co before 17:00
                if co and co.hour*60+co.minute < 17*60:
                    if status=='present':
                        # mark early_leave if very early? keep present for now
                        pass
            # working hours
            wh = None
            ot = 0.0
            if ci and co:
                diff = (co.hour*60+co.minute) - (ci.hour*60+ci.minute)
                wh = max(0, diff/60.0 - 1.0)  # minus 1h break
                if co.hour*60+co.minute > 17*60+15:
                    ot = (co.hour*60+co.minute - (17*60+15))/60.0
            now_ts = int(datetime.datetime.now(tz=datetime.timezone.utc).timestamp())
            if existing:
                # update
                cur.execute("UPDATE attendance_table SET check_in_time=?, check_out_time=?, working_hours=?, overtime_hours=?, status=?, late_minutes=?, permission_hours=?, updated_at=? WHERE id=?",
                            (ci_ts, co_ts, wh, ot, status, 0, 0.0, now_ts, existing[0]))
                updated_att +=1
            else:
                cur.execute("INSERT INTO attendance_table (staff_id, date, check_in_time, check_out_time, working_hours, overtime_hours, status, late_minutes, permission_hours, excused, excused_hours, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                            (sid, date_ts, ci_ts, co_ts, wh, ot, status, 0, 0.0, 0, 0.0, now_ts, now_ts))
                inserted_att +=1
    # monthly summary
    inserted_month=0
    updated_month=0
    for sid, vals in monthly_from_file.items():
        period = vals['period']
        cur.execute("SELECT id FROM monthly_attendance_summary_table WHERE staff_id=? AND period=?", (sid, period))
        ex = cur.fetchone()
        now_ts2 = int(datetime.datetime.now(tz=datetime.timezone.utc).timestamp())
        # monthly table uses text period and real hours, int absent
        if ex:
            cur.execute("UPDATE monthly_attendance_summary_table SET late_hours=?, overtime_hours=?, excused_hours=?, absent_days=?, updated_at=? WHERE id=?",
                        (vals['late_hours'], vals['overtime_hours'], vals['permission_hours'], vals['absent_days'], now_ts2, ex[0]))
            updated_month+=1
        else:
            cur.execute("INSERT INTO monthly_attendance_summary_table (staff_id, period, late_hours, overtime_hours, excused_hours, absent_days, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?)",
                        (sid, period, vals['late_hours'], vals['overtime_hours'], vals['permission_hours'], vals['absent_days'], now_ts2, now_ts2))
            inserted_month+=1
    con.commit()
    print(f"ATT inserted {inserted_att} updated {updated_att}")
    print(f"MONTH inserted {inserted_month} updated {updated_month}")
    # verify
    cur.execute("SELECT COUNT(*) FROM attendance_table WHERE date BETWEEN ? AND ?", (int(datetime.datetime(2026,8,1,tzinfo=datetime.timezone.utc).timestamp()), int(datetime.datetime(2026,8,31,tzinfo=datetime.timezone.utc).timestamp())))
    print("attendance Aug count", cur.fetchone())
    cur.execute("SELECT COUNT(*) FROM monthly_attendance_summary_table WHERE period='2026-08'")
    print("monthly 2026-08 count", cur.fetchone())
except Exception as e:
    con.rollback()
    print("ERROR", e)
    import traceback; traceback.print_exc()
finally:
    con.close()
print("DONE")
